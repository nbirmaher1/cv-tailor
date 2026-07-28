#!/usr/bin/env python3
"""Render a tailored cover letter from a simple content JSON file into a .docx.

JSON schema:
{
  "full_name": str, "contact_line": str, "date": str, "salutation": str,
  "paragraphs": [str], "closing": str, "signature_name": str
}
"""
import json
import sys

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt


def build(data, output_path):
    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal_style.paragraph_format.line_spacing = 1.4

    name_p = doc.add_paragraph()
    name_run = name_p.add_run(data["full_name"])
    name_run.bold = True
    name_run.font.size = Pt(15)

    if data.get("contact_line"):
        cp = doc.add_paragraph(data["contact_line"])
        for run in cp.runs:
            run.font.size = Pt(9.5)

    date_p = doc.add_paragraph(data.get("date", ""))
    date_p.paragraph_format.space_before = Pt(18)
    date_p.paragraph_format.space_after = Pt(14)

    sal_p = doc.add_paragraph(data.get("salutation", ""))
    sal_p.paragraph_format.space_after = Pt(10)

    for para in data.get("paragraphs", []):
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(10)

    closing_p = doc.add_paragraph(data.get("closing", ""))
    closing_p.paragraph_format.space_before = Pt(10)

    if data.get("signature_name"):
        sig_p = doc.add_paragraph(data["signature_name"])
        sig_run = sig_p.runs[0]
        sig_run.bold = True

    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: render_cover_letter_docx.py <content.json> <output.docx>", file=sys.stderr)
        sys.exit(1)
    with open(sys.argv[1]) as f:
        content = json.load(f)
    build(content, sys.argv[2])
    print(f"Wrote {sys.argv[2]}")
